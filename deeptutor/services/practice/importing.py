"""Bounded, non-executing parsers for the learner's CSV, XLSX and JSON files."""

from __future__ import annotations

import csv
import hashlib
import io
import json
from pathlib import Path
import re
import zipfile

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 500
MAX_CELL = 20000
MAX_COLUMNS = 64
ALIASES = {
    "question": ("question", "题目", "题干"),
    "question_type": ("question_type", "type", "题型"),
    "correct_answer": ("correct_answer", "answer", "答案", "正确答案"),
    "explanation": ("explanation", "解析"),
    "difficulty": ("difficulty", "难度"),
    "tags": ("tags", "标签", "分类"),
    "user_answer": ("user_answer", "我的答案", "作答"),
    "options": ("options", "选项"),
}
TYPES = {
    "single_choice": "single_choice",
    "multiple_choice": "single_choice",
    "mcq": "single_choice",
    "单选": "single_choice",
    "单选题": "single_choice",
    "multi_choice": "multi_choice",
    "多选": "multi_choice",
    "多选题": "multi_choice",
    "true_false": "true_false",
    "判断": "true_false",
    "判断题": "true_false",
    "fill_blank": "fill_blank",
    "填空": "fill_blank",
    "填空题": "fill_blank",
    "short_answer": "short_answer",
    "简答": "short_answer",
    "简答题": "short_answer",
    "essay": "short_answer",
    "free_response": "short_answer",
}


def _cell(value: object) -> str:
    text = "" if value is None else str(value).strip()
    if len(text) > MAX_CELL:
        raise ValueError("A cell exceeds 20,000 characters")
    return text


def parse_rows(data: bytes, filename: str) -> list[tuple[int, dict]]:
    if not data or len(data) > MAX_BYTES:
        raise ValueError("Choose a non-empty file up to 5 MB")
    suffix = Path(filename).suffix.lower()
    if suffix == ".json":
        rows = json.loads(data.decode("utf-8-sig"))
        if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
            raise ValueError("JSON must contain an array of question objects")
        if len(rows) > MAX_ROWS:
            raise ValueError("Import at most 500 questions at a time")
        return list(enumerate(rows, 1))
    if suffix in {".csv", ".tsv"}:
        try:
            content = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = data.decode("gb18030")
        try:
            dialect = csv.Sniffer().sniff(content[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
        rows = []
        try:
            for row in csv.reader(io.StringIO(content, newline=""), dialect):
                if len(row) > MAX_COLUMNS:
                    raise ValueError("Import at most 64 columns at a time")
                rows.append(row)
                if len(rows) > MAX_ROWS + 1:
                    raise ValueError("Import at most 500 questions at a time")
        except csv.Error as exc:
            raise ValueError("Invalid CSV data or a cell exceeds the supported size") from exc
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError("The expanded workbook exceeds 30 MB")
            if len(archive.infolist()) > 2000:
                raise ValueError("The workbook contains too many parts")
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = book.active
            if sheet is None:
                raise ValueError("The workbook has no active worksheet")
            if (sheet.max_column or 0) > MAX_COLUMNS:
                raise ValueError("Import at most 64 columns at a time")
            rows = []
            for cells in sheet.iter_rows():
                if len(cells) > MAX_COLUMNS:
                    raise ValueError("Import at most 64 columns at a time")
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("Replace spreadsheet formulas with values before importing")
                rows.append([cell.value for cell in cells])
                if len(rows) > MAX_ROWS + 1:
                    raise ValueError("Import at most 500 questions at a time")
        finally:
            book.close()
    else:
        raise ValueError("Supported formats: .xlsx, .csv, .tsv and .json")
    if not rows:
        raise ValueError("The file is empty")
    headers = [_cell(value).lower() for value in rows[0]]
    nonempty = [header for header in headers if header]
    if len(nonempty) != len(set(nonempty)):
        raise ValueError("Column names must be unique")
    if len(rows) > MAX_ROWS + 1:
        raise ValueError("Import at most 500 questions at a time")
    return [
        (n, dict(zip(headers, row, strict=False)))
        for n, row in enumerate(rows[1:], 2)
        if any(_cell(value) for value in row)
    ]


def normalize_question(raw: dict) -> dict:
    row = {str(key).strip().lower(): value for key, value in raw.items()}
    fields = {key: next((row[a] for a in names if a in row), "") for key, names in ALIASES.items()}
    question, answer = _cell(fields["question"]), _cell(fields["correct_answer"])
    if not question or not answer:
        raise ValueError("Question and correct answer are required")
    options = fields["options"] or {}
    if isinstance(options, str):
        options = json.loads(options)
    if isinstance(options, list):
        options = {chr(65 + i): value for i, value in enumerate(options)}
    if not isinstance(options, dict) or len(options) > 10:
        raise ValueError("Options must contain at most 10 named choices")
    options = {
        str(key).strip().upper(): _cell(value) for key, value in options.items() if _cell(value)
    }
    for letter in "ABCDEFGHIJ":
        value = row.get(letter.lower(), row.get(f"选项{letter.lower()}", ""))
        if _cell(value):
            options[letter] = _cell(value)
    raw_type = _cell(fields["question_type"]).lower()
    kind = TYPES.get(raw_type) if raw_type else ("single_choice" if options else "short_answer")
    if kind is None:
        raise ValueError("Unknown question type")
    if kind == "true_false":
        truths = {"true", "正确", "对", "是", "t", "1"}
        falses = {"false", "错误", "错", "否", "f", "0"}
        if answer.lower() not in truths | falses:
            raise ValueError("A true/false answer must be true or false")
        options = {"T": "True", "F": "False"}
        answer = "T" if answer.lower() in truths else "F"
    if kind in {"single_choice", "multi_choice"}:
        if len(options) < 2:
            raise ValueError("Choice questions need at least two options")
        if answer in options.values():
            answer = next(key for key, text in options.items() if text == answer)
        keys = [key.strip().upper() for key in re.split(r"[,;，；\s]+", answer) if key.strip()]
        if kind == "multi_choice" and len(keys) == 1 and keys[0] not in options:
            keys = list(keys[0])
        if (
            not keys
            or any(key not in options for key in keys)
            or (kind == "single_choice" and len(keys) != 1)
        ):
            raise ValueError("The answer must name existing option keys (multi-choice: A,C)")
        answer = ",".join(sorted(set(keys)))
    tags_value = fields["tags"]
    tags = tags_value if isinstance(tags_value, list) else re.split(r"[,;，；]", _cell(tags_value))
    tags = list(dict.fromkeys(_cell(tag) for tag in tags if _cell(tag)))
    if len(tags) > 20 or any(len(tag) > 100 for tag in tags):
        raise ValueError("Use at most 20 tags, each up to 100 characters")
    result = dict(
        question=question,
        question_type=kind,
        options=options,
        correct_answer=answer,
        explanation=_cell(fields["explanation"]),
        difficulty=_cell(fields["difficulty"]),
        user_answer=_cell(fields["user_answer"]),
        tags=tags,
    )
    identity = {
        key: result[key] for key in ("question", "question_type", "options", "correct_answer")
    }
    result["question_id"] = (
        "import:"
        + hashlib.sha256(
            json.dumps(identity, sort_keys=True, ensure_ascii=False).encode()
        ).hexdigest()
    )
    return result


def preview(data: bytes, filename: str) -> dict:
    questions, errors = [], []
    for number, row in parse_rows(data, filename):
        try:
            questions.append(normalize_question(row))
        except (ValueError, TypeError) as exc:
            errors.append({"row": number, "message": str(exc)})
    if not questions and not errors:
        raise ValueError("The file contains no questions")
    return {"questions": questions, "errors": errors}
